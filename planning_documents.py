"""Create editable Word and review PDF versions of an activity plan."""

from __future__ import annotations

import uuid
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parent
EXPORT_DIR = ROOT / "data" / "activity_exports"
FONT_NAME = "COEChinese"


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(_text(item) for item in value if item is not None)
    if isinstance(value, dict):
        return "；".join(f"{key}：{_text(item)}" for key, item in value.items())
    return str(value).strip()


def _sections(plan: dict) -> list[dict]:
    sections = plan.get("document_sections")
    if isinstance(sections, list) and sections:
        clean = []
        for section in sections:
            if isinstance(section, dict):
                clean.append({"heading": _text(section.get("heading")), "body": _text(section.get("body")), "table": section.get("table")})
        if clean:
            return clean
    mapping = [
        ("活动背景与目标", ("objective", "key_message")),
        ("活动对象与概要", ("audience",)),
        ("筹备与时间安排", ("schedule",)),
        ("具体流程", ("workflow",)),
        ("宣传与物资", ("promotion_plan", "materials")),
        ("资源与注意事项", ("budget_or_resources", "risk_controls", "open_questions")),
    ]
    return [{"heading": heading, "body": "\n".join(_text(plan.get(key)) for key in keys if plan.get(key))}
            for heading, keys in mapping if any(plan.get(key) for key in keys)]


def _table_parts(table: object) -> tuple[list[str], list[list[str]]]:
    if not isinstance(table, dict):
        return [], []
    headers = [_text(value) for value in table.get("headers", [])]
    rows = [[_text(value) for value in row] for row in table.get("rows", []) if isinstance(row, list)]
    if not headers or not rows:
        return [], []
    width = len(headers)
    return headers, [(row + [""] * width)[:width] for row in rows]


def _set_east_asia(run, font_name: str = "Microsoft YaHei") -> None:
    run.font.name = font_name
    run.font.size = Pt(10.5)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), font_name)


def _write_docx(plan: dict, path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.9)
    section.left_margin = section.right_margin = Cm(2.4)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.22
    for name, size in (("Title", 18), ("Heading 1", 13)):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(9 if name == "Heading 1" else 0)
        style.paragraph_format.space_after = Pt(5)
    # Word's built-in Title style can inherit a blue bottom rule from the
    # default template; explicitly remove it from the exported document.
    title_ppr = document.styles["Title"]._element.get_or_add_pPr()
    borders = title_ppr.find(qn("w:pBdr"))
    if borders is not None:
        title_ppr.remove(borders)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "nil")
    borders.append(bottom)
    title_ppr.append(borders)
    title = _text(plan.get("title")) or "活动策划案"
    document.add_paragraph(title, "Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    for section_data in _sections(plan):
        heading = section_data["heading"] or "活动安排"
        document.add_paragraph(heading, "Heading 1")
        body = section_data["body"]
        for line in body.splitlines():
            if line.strip():
                paragraph = document.add_paragraph()
                _set_east_asia(paragraph.add_run(line.strip()))
        headers, rows = _table_parts(section_data.get("table"))
        if headers:
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            table.autofit = False
            widths = [Cm(2.3), Cm(3.0), Cm(10.2)] if len(headers) == 3 else [Cm(15.5 / len(headers))] * len(headers)
            for column, width in zip(table.columns, widths):
                column.width = width
            for col, value in enumerate(headers):
                _set_east_asia(table.rows[0].cells[col].paragraphs[0].add_run(value))
            for row in rows:
                cells = table.add_row().cells
                for col, value in enumerate(row):
                    _set_east_asia(cells[col].paragraphs[0].add_run(value))
            for table_row in table.rows:
                for cell, width in zip(table_row.cells, widths):
                    cell.width = width
            document.add_paragraph().paragraph_format.space_after = Pt(0)
    questions = plan.get("open_questions")
    if questions and not any("待确认" in section["heading"] for section in _sections(plan)):
        document.add_paragraph("待确认事项", "Heading 1")
        for line in _text(questions).splitlines():
            document.add_paragraph(line)
    document.save(path)


def _register_font() -> str:
    if FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        for candidate in (Path("C:/Windows/Fonts/simhei.ttf"), Path("C:/Windows/Fonts/Deng.ttf")):
            if candidate.is_file():
                pdfmetrics.registerFont(TTFont(FONT_NAME, str(candidate)))
                break
        else:
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            return "STSong-Light"
    return FONT_NAME


def _pdf_paragraph(value: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(value).replace("\n", "<br/>"), style)


def _write_pdf(plan: dict, path: Path) -> None:
    font = _register_font()
    title_style = ParagraphStyle("title", fontName=font, fontSize=17, leading=26, alignment=TA_CENTER, spaceAfter=20, wordWrap="CJK", textColor=colors.black)
    heading_style = ParagraphStyle("heading", fontName=font, fontSize=12, leading=19, spaceBefore=12, spaceAfter=6, wordWrap="CJK", textColor=colors.black)
    body_style = ParagraphStyle("body", fontName=font, fontSize=10, leading=17, spaceAfter=6, alignment=TA_LEFT, wordWrap="CJK", textColor=colors.black)
    cell_style = ParagraphStyle("cell", parent=body_style, fontSize=9, leading=14, spaceAfter=0)
    story = [_pdf_paragraph(_text(plan.get("title")) or "活动策划案", title_style)]
    for section in _sections(plan):
        story.append(_pdf_paragraph(section["heading"] or "活动安排", heading_style))
        for line in section["body"].splitlines():
            if line.strip():
                story.append(_pdf_paragraph(line.strip(), body_style))
        headers, rows = _table_parts(section.get("table"))
        if headers:
            data = [[_pdf_paragraph(value, cell_style) for value in row] for row in [headers, *rows]]
            available = A4[0] - 2 * 65
            if len(headers) == 3:
                widths = [available * .17, available * .2, available * .63]
            else:
                widths = [available / len(headers)] * len(headers)
            table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), .45, colors.HexColor("#d9d9d9")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8edf5")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]))
            story.extend([Spacer(1, 5), table, Spacer(1, 7)])
    questions = plan.get("open_questions")
    if questions and not any("待确认" in section["heading"] for section in _sections(plan)):
        story.append(_pdf_paragraph("待确认事项", heading_style))
        for line in _text(questions).splitlines():
            story.append(_pdf_paragraph(line, body_style))
    doc = SimpleDocTemplate(str(path), pagesize=A4, topMargin=58, bottomMargin=58, leftMargin=65, rightMargin=65)
    def footer(canvas, doc):
        canvas.setFont(font, 9)
        canvas.drawCentredString(A4[0] / 2, 35, f"第 {doc.page} 页")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def write_documents(plan: dict, project_id: str) -> dict:
    project_id = str(uuid.UUID(project_id))
    folder = EXPORT_DIR / project_id
    folder.mkdir(parents=True, exist_ok=True)
    docx_path, pdf_path = folder / "plan.docx", folder / "plan.pdf"
    _write_docx(plan, docx_path)
    _write_pdf(plan, pdf_path)
    return {"docx_url": f"/api/activity/documents/{project_id}/docx", "pdf_url": f"/api/activity/documents/{project_id}/pdf"}


def document_path(project_id: str, format_name: str) -> Path:
    project_id = str(uuid.UUID(project_id))
    if format_name not in {"docx", "pdf"}:
        raise ValueError("不支持的文档格式")
    return EXPORT_DIR / project_id / f"plan.{format_name}"
